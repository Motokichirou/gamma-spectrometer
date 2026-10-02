# -*- coding: utf-8 -*-
"""Сборка stock/bom_zakupka.xlsx из pcb/gamma_spectrometer.net.

Два листа, вместе покрывают все 178 компонентов:
  «К заказу»   — только то, что действительно покупаем, с фактическим требованием
                 (не PN из BOM там, где PN из BOM не годится).
  «Из запасов» — что ставим своё, с указанием, вместо чего по BOM.

Ссылки на магазин берутся из пользовательского поля схемы («Ссылка на закупку»).
Решения по заменам живут в таблице PLAN ниже (сверка с stock/required.md,
фото номенклатуры 2026-09-23). При изменении запасов править PLAN, а не .xlsx.

Запуск: python stock/make_bom_xlsx.py
"""
import io
import os
import re
import sys
from collections import OrderedDict

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NET = os.path.join(ROOT, 'pcb', 'gamma_spectrometer.net')
OUT = os.path.join(ROOT, 'stock', 'bom_zakupka.xlsx')

BUY, HAVE, NA = 'buy', 'have', 'na'

# Ключ: PN из BOM (или номинал, если PN не задан).
#   BUY  -> ('buy',  что заказать, номинал, корпус, сколько заказать, примечание)
#   HAVE -> ('have', что ставим,   номинал, корпус, вместо-чего-пояснение)
#   NA   -> ('na',   пояснение)
# «Номинал/корпус» — ФАКТИЧЕСКИЕ, то есть то, что реально встанет на плату.
PLAN = {
    # ================= ЗАКАЗАТЬ =================
    'CC0603KRX7R9BB103': (BUY, 'CC0603KRX7R0BB103', '10nF 100 В', '0603', 5,
                          'ЗАКАЗАНО 2026-10-02. PN из BOM (код 9BB) на 50 В не годится: снаббер сидит на '
                          '35.7 В с пиками 43-45 В. Взят 0BB = 100 В, та же серия Yageo, X7R, +-10 %. '
                          'Корпус строго 0603 — 0805 на пад не сядет'),
    'CC0603FRNPO9BN511': (BUY, 'CC0603FRNPO9BN511', '510pF', '0603', 5,
                          'C0G, допуск +-1 % — интегратор шейпера, задаёт форму импульса'),
    'CC1206JKNPOCBN221': (BUY, 'CC1206JKNPOCBN221', '220pF', '1206', 5,
                          'C0G, 1 кВ — ВЧ-фильтр выхода ВВ, два в серии'),
    'CC1210KKX7RCBB223': (BUY, 'CC1210KKX7RCBB223', '22nF', '1210', 3,
                          '>=250 В — накопитель между динодами ФЭУ'),
    'GRM43DR73A473KW01L': (BUY, 'GRM43DR73A473KW01L', '47nF', '1812', 3,
                           '>=250 В — накопитель между динодами ФЭУ'),
    'RC1206FR-071M6L': (BUY, 'RC1206FR-071M6L', '1.6M', '1206', 5,
                        '155 В при уставке 1200 В — выше предела 0805 (150 В), нужен 1206'),
    'RC1206FR-072ML': (BUY, 'RC1206FR-072ML', '2M', '1206', 5,
                       '193 В при 1200 В — только 1206. При 1500 В будет 242 В — уставку не задирать'),
    'ABM8G-16.000MHZ-4Y-T3': (BUY, 'ABM8G-16.000MHZ-4Y-T3', '16 МГц', '3225, 4 вывода', 2,
                              'заменяется легко, но CL должен остаться ~8-9 пФ (у нас C301/C302 = 18 пФ)'),
    '1N4148W': (BUY, '1N4148W', '1N4148W', 'SOD-123', 10,
                'ССЫЛКИ В СХЕМЕ НЕТ — дозаполнить поле. Защита баз буферов динодов'),
    'MF-MSMF050-2': (BUY, 'MF-MSMF050-2', '0.5 A', '1812', 2,
                     'самовосстанавливающийся предохранитель по +5V_USB'),
    'STM32G474MET6': (BUY, 'STM32G474M[B/C/E]T6, G484M*, G473M*, G483M*', 'LQFP-80', 'LQFP-80 12x12', 2,
                      'ВНИМАНИЕ: буква M = LQFP-80 обязательна, RET6 (64) и CET6 (48) на плату не встанут. '
                      'Годится любой флеш B/C/E (128/256/512 КБ — прошивка влезает в 128). Годятся G484 (тот же + AES), '
                      'G473 и G483 (минус HRTIM, мы его не используем; ADC4 и 6 ОУ на месте). '
                      'НЕ годится G471/G491 — там 3 АЦП, нет ADC4 под Pt1000 на PB12. '
                      'Купленный чип проверить по SWD: PKG 0x1FFF7500 = 0x05, FLASH_SIZE 0x1FFF75E0, DEV_ID 0x469'),
    'BLM18EG121SN1D': (BUY, 'BLM18EG121SN1D', '120 Ом @100 МГц', '0603', 3,
                       'феррит по VDDA; заменяется любым 120 Ом @100 МГц в 0603'),
    'SRU5028-101Y': (BUY, 'SRU5028-101Y', '100uH', 'SRU5028 5.2x5.2', 2,
                     'через L103 идёт вся ВВ-плата — смотреть ток насыщения, не только индуктивность'),
    'SRN3015-470M': (BUY, 'SRN3015-470M', '47uH', 'SMS-ME3015', 3,
                     'LC-фильтры +-5V_A, ток небольшой — заменяется по габариту'),
    'FT231XQ': (BUY, 'FT231XQ', 'FT231XQ', 'QFN-20 4x4', 2,
                'ВНИМАНИЕ: FTDI массово подделывают, поддельные отваливаются после обновления драйвера — брать у дистрибьютора'),
    'ADA4817-1ARDZ': (BUY, 'ADA4817-1ARDZ', 'ADA4817-1ARDZ', 'SOIC-8 EP', 2,
                      'ССЫЛКИ НЕТ, гуглить отдельно. Именно -1ARDZ (SOIC-8 с площадкой); ACPZ — это LFCSP, не встанет. Замены в схеме нет, брать 2 шт'),
    'W25Q64JVSSIQ': (BUY, 'W25Q64JVSSIQ', '64 Мбит', 'SOIC-8', 2, 'SPI flash под спектры и калибровку'),
    'PRTR5V0U2X': (BUY, 'PRTR5V0U2X', 'PRTR5V0U2X', 'SOT-143', 3, 'ESD-защита линий USB'),
    'AMS1117-3.3': (BUY, 'AMS1117-3.3', '3.3 В', 'SOT-223', 3,
                    'на выходе обязателен тантал (есть в запасах) — с керамикой регулятор неустойчив'),
    'BAV23S': (BUY, 'BAV23S', 'BAV23S', 'SOT-23', 12,
               'диоды умножителя, 500 В на корпус; вывод 3 не подключён'),
    'IRLML0060': (BUY, 'IRLML0060', '60 В', 'SOT-23', 3,
                  'ВНИМАНИЕ: именно 0060. IRLML0040 в запасах есть — ставить НЕЛЬЗЯ: пик на стоке 43-45 В'),
    'MMBTA42': (BUY, 'MMBTA42', '300 В', 'SOT-23', 6, 'буферы динодов Dy6-Dy8'),
    'TL431': (BUY, 'TL431', '2.5 В', 'SOT-23', 3,
              'опора монитора ВВ; точность влияет только на шкалу монитора и калибруется'),
    'TPS60400DBV': (BUY, 'TPS60400DBV', '-5 В', 'SOT-23-5', 2,
                    'C110/C111 к нему обязательно X7R (есть в запасах)'),
    'REF35250QDBVR': (BUY, 'REF35250QDBVR', '2.500 В', 'SOT-23-6', 2,
                      'ССЫЛКИ НЕТ, гуглить отдельно. Сверить по даташиту: последовательная опора 2.5 В в SOT-23-6. TL431 не замена — это точность всей шкалы'),
    'DX07S024WJ3R400': (BUY, 'DX07S024WJ3R400', 'USB-C', 'JAE DX07S024', 2,
                        'футпринт под конкретный корпус JAE — замена означает переразводку'),

    # ================= ИЗ ЗАПАСОВ =================
    'TR3B106K010C0750': (HAVE, 'CA45-A016K106T (тантал)', '10uF 16 В', 'корпус A (EIA-3216)',
                         'корпус A меньше пада 3528 — встаёт. 20 шт. Для C105 тантал обязателен: AMS1117 с керамикой неустойчив'),
    'CC0402FRNPO9BN180': (HAVE, '18 пФ C0G из каталога 0402', '18pF 50 В', '0402',
                          'нагрузочные кварца, C0G +-5 %'),
    'GRM1555C1H220JA01D': (HAVE, '22 пФ C0G из каталога 0402', '22pF 50 В', '0402',
                           'C_COMP2 на COMP MAX1847'),
    'GRM1555C1H4R7WA01D': (HAVE, '4.7 пФ C0G из каталога 0402', '4.7pF 50 В', '0402',
                           'допуск +-0.25 пФ вместо +-0.05 -> полоса TIA гуляет на +-5 %. Приемлемо'),
    'CC0603JRX7R9BB104': (HAVE, 'GRM155R71H104KE14D', '0.1uF 50 В X7R', '0402',
                          '54 шт; 50 В вместо 16 В из каталога. Все позиции на шинах <=5 В'),
    'CC0603JRNPO9BN101': (HAVE, 'CC0402FRNPO9BN101', '100pF 50 В C0G +-1 %', '0402',
                          'допуск как в BOM — для C315 (звено CR шейпера) это важно'),
    'CC0603JRX7R7BB105': (HAVE, 'GRM155R61C105KA12D', '1uF 16 В X5R', '0402',
                          'разделительные в тракте'),
    'CC0805JKX7R9BB474': (HAVE, 'GRM155R61C105KA12D', '1uF 16 В X5R (вместо 0.47uF)', '0402',
                          'под 4.25 В останется ~0.7 мкФ — ближе к даташиту MAX1847, чем 0.47'),
    'CC0805JKNPOBBN102': (HAVE, 'CC0805JKNPOBBN102', '1nF 500 В C0G', '0805',
                          'ровно этот PN, 30 шт. Ступени умножителя, размах 310-385 В'),
    'CC0805JKX7R9BB105': (HAVE, 'GRM155R61C105KA12D', '1uF 16 В X5R', '0402',
                          'на опоре 2.5 В останется ~0.8 мкФ, REF35250 хватает (нужно 0.1-10 мкФ)'),
    'CC0805JRX7R9BB333': (HAVE, 'GRM155R71H333KE14D', '33nF 50 В X7R', '0402',
                          'цепь коррекции петли ВВ, на COMP всего 1.3 В'),
    'GRM32ER71A476KE15L': (HAVE, '22 мкФ 35 В X5R из запасов', '22uF 35 В (вместо 47uF)', '1206',
                           'под 5 В даёт ~20 мкФ против расчётных 25 — модель это терпит, пульсации чуть выше'),
    'C1812X104K102T': (HAVE, 'C1812X104K102T', '0.1uF 1000 В', '1812',
                       'ровно этот PN, 12 шт при потребности 11 — запасной всего один, паять аккуратно'),
    'RC0402FR-0722RL': (HAVE, '22 Ом из книги 0402', '22 Ом 1 %', '0402',
                        'R_SNUB на выходе TIA — без него TIA генерирует'),
    'RL0603FR-070R1L': (HAVE, '0.1 Ом 1 % отдельной лентой', '0.1 Ом 1 %', '0603',
                        'датчик тока MAX1847; 16 мВт при импульсах 0.8 А — 0603 хватает'),
    'RC0603FR-071K2L': (HAVE, 'из книги 0603', '1.2k 1 %', '0603', ''),
    'RC0603FR-07100KL': (HAVE, 'из книги 0603', '100k 1 %', '0603', 'на каждом ~20 В (ток 216 мкА)'),
    'AC0603FR-0710RL': (HAVE, 'из книги 0603', '10 Ом 1 %', '0603', ''),
    'RC0603FR-0710KL': (HAVE, 'из книги 0603', '10k 1 %', '0603', ''),
    'RC0603FR-07110KL': (HAVE, 'из книги 0603', '110k 1 %', '0603', 'нижнее плечо монитора ВВ'),
    'RC0603FR-0714K7L': (HAVE, 'из книги 0603', '15k 1 % (вместо 14.7k)', '0603',
                         '14.7k — ряд E96. Смещение АЦП уедет на 2 %, но базовую линию DSP измеряет сам'),
    'RC0603FR-07180KL': (HAVE, 'из книги 0603', '180k 1 %', '0603', 'задаёт частоту MAX1847 252 кГц'),
    'RT0603BRD071ML': (HAVE, 'из книги 0603', '1M 1 %', '0603', 'цепь коррекции COMP'),
    'RC0603FR-071KL': (HAVE, 'из книги 0603', '1k 1 %', '0603', ''),
    'RC0603FR-072M2L': (HAVE, 'из книги 0603', '2.2M 1 %', '0603',
                        'нижнее плечо FB, задаёт коэффициент ВВ x454.5'),
    'RC0603FR-0722KL': (HAVE, 'из книги 0603', '22k 1 %', '0603', ''),
    'RT0603BRD07330RL': (HAVE, 'из книги 0603', '330 Ом 1 %', '0603', ''),
    'RC1206FR-07499RL': (HAVE, 'из книги 0603', '510 Ом 1 % (вместо 499)', '0603',
                         'ВНИМАНИЕ: в схеме PN и ссылка от 1206, а посадка 0603 — поправить поле. '
                         '510 вместо 499: усиление TIA +2.2 %, энергокалибровка поглотит'),
    'RC0603FR-075K1L': (HAVE, 'из книги 0603', '5.1k 1 %', '0603', 'CC-резисторы USB-C'),
    'RC0603FR-078K2L': (HAVE, 'из книги 0603', '8.2k 1 %', '0603', ''),
    'RC0603FR-07909RL': (HAVE, 'из книги 0603', '910 Ом 1 % (вместо 909)', '0603',
                         '909 — ряд E96; отличие 0.1 %, на смещение АЦП не влияет'),
    'RC1206FR-071M3L': (HAVE, 'из книги 0805', '1.3M 1 %', '0805 (вместо 1206)',
                        '126 В при уставке 1200 В против предела 0805 в 150 В'),
    'RC1206JR-07100ML': (HAVE, 'RC0805JR-07100ML', '100M 5 %', '0805 (вместо 1206)',
                         '50 шт. 120 В на резистор при 1200 В, предел 0805 = 150 В — уставку выше не задирать'),
    'RC1206FR-076M8L': (HAVE, 'из книги 0805', '6.8M 1 %', '0805 (вместо 1206)',
                        'делитель монитора ВВ, 120 В при 1200 В'),
    'RC1206FR-07680KL': (HAVE, 'из книги 0805', '680k 1 %', '0805 (вместо 1206)',
                         'делитель ФЭУ, 66 В при 1200 В'),
    'RC1206FR-07750KL': (HAVE, 'из книги 0805', '750k 1 %', '0805 (вместо 1206)',
                         'делитель ФЭУ, 73 В при 1200 В'),
    'max1847eee+': (HAVE, 'MAX1847EEE+', 'MAX1847EEE+', 'QSOP-16',
                    '3 шт в наличии — самая дорогая позиция проекта'),
    'ATB322524-0110-T000': (HAVE, 'ATB322524-0110-T000', '1:10.2', 'SMD 3.2x2.5',
                            'тот самый экземпляр, под который считался узел ВВ. Запасного нет — паять аккуратно, феррит перегрева не прощает'),
    'BZT52C18': (HAVE, 'BZT52C18', '18 В', 'SOD-123',
                 '50 шт. ВНИМАНИЕ: проверить корпус SOD-123 (не -123W/F). Отобрать пару с близким напряжением — они задают пик на стоке'),
    'BAS516': (HAVE, 'BAS516 (маркировка YJ)', 'BAS516', 'SOD-523',
               '33 шт. Диод снаббера — именно BAS516, BAS16J в SOD-323 на пад не встанет'),
    'FYLS-0805UBC': (HAVE, 'LTST-C195KGJRKT', 'зелёный', '0603 на пад 0805', 'светодиод индикации'),
    '150080GS75000': (HAVE, 'LTST-C195KGJRKT', 'зелёный', '0603 на пад 0805', 'светодиод индикации'),
    'LP5907MFX-3.3': (HAVE, 'LP2985AIM5X-3.3', '3.3 В', 'SOT-23-5',
                      '10 шт. Вывод 4 у нас не подключён, а у LP2985 это BYPASS — оставить в воздухе. '
                      'Шум 30 мкВ против 6.5, но эта шина питает опору и VDDA — влияние ничтожно'),

    # ================= НИ ТО НИ ДРУГОЕ =================
    'PMT': (NA, 'паяется проводами напрямую к выводам ФЭУ R1307-01 — покупать нечего'),
}

# Позиции, которые делятся между «заказать» и «из запасов».
SPLIT = {
    'GRM21BR61C475KA88L': dict(
        buy_refs=['C411', 'C413'],
        buy=('4.7 мкФ X5R/X7R, >=16 В', '4.7uF >=16 В', '0805', 5,
             'питание ADA4817 по +-5 В: своё 0402 на 6.3 В даёт запас всего 1.26x и теряет ~70 % ёмкости'),
        have=('4.7 мкФ 6.3 В из каталога 0402', '4.7uF 6.3 В', '0402',
              'C106/C310 сидят на 3.3 В — запаса хватает'),
    ),
}

SHEET_RU = {'USB_Power_Board': 'USB/питание', 'HV_Board': 'ВВ',
            'MCU_Board': 'МК', 'Divider_Board': 'делитель', '': '—'}


def parse_net():
    s = io.open(NET, encoding='utf-8').read()
    body = s[s.index('(components'):s.index('(libparts')]
    comps = []
    for m in re.finditer(r'\(comp \(ref "([^"]+)"\)(.*?)\n      \(tstamps ', body, re.S):
        ref, blk = m.group(1), m.group(2)

        def one(p):
            r = re.search(p, blk)
            return r.group(1) if r else ''

        pn, link = '', ''
        for fm in re.finditer(r'\(field \(name "([^"]*)"\) "([^"]*)"\)', blk):
            nm, v = fm.group(1).strip(), fm.group(2).strip()
            if nm.upper() in ('PN', 'MPN'):
                pn = v
            elif v.startswith('http') and nm.lower() != 'datasheet':
                link = v
        comps.append(dict(ref=ref, val=one(r'\(value "([^"]*)"\)'), pn=pn, link=link,
                          sheet=SHEET_RU.get(one(r'\(property \(name "Sheetname"\) \(value "([^"]*)"\)\)'), '?')))
    return comps


def refkey(r):
    m = re.match(r'([A-Za-z_]+)(\d*)', r)
    return (m.group(1), int(m.group(2) or 0))


def fmt(refs, comps_by_ref):
    refs = sorted(refs, key=refkey)
    sheets = sorted({comps_by_ref[r]['sheet'] for r in refs})
    return ', '.join(refs), ', '.join(sheets)


def main():
    comps = parse_net()
    by_ref = {c['ref']: c for c in comps}
    groups = OrderedDict()
    for c in comps:
        groups.setdefault(c['pn'] or c['val'], dict(refs=[], link=''))
        groups[c['pn'] or c['val']]['refs'].append(c['ref'])
        if c['link']:
            groups[c['pn'] or c['val']]['link'] = c['link']

    buy, have, na, unknown = [], [], [], []
    for key, g in groups.items():
        if key in SPLIT:
            sp = SPLIT[key]
            br = [r for r in g['refs'] if r in sp['buy_refs']]
            hr = [r for r in g['refs'] if r not in sp['buy_refs']]
            what, val, fp, oq, note = sp['buy']
            buy.append(dict(what=what, val=val, fp=fp, oq=oq, note=note,
                            need=len(br), refs=br, link=g['link'], bom=key))
            what, val, fp, note = sp['have']
            have.append(dict(what=what, val=val, fp=fp, note=note, refs=hr, bom=key))
            continue
        if key not in PLAN:
            unknown.append(key)
            continue
        p = PLAN[key]
        if p[0] == BUY:
            _, what, val, fp, oq, note = p
            buy.append(dict(what=what, val=val, fp=fp, oq=oq, note=note,
                            need=len(g['refs']), refs=g['refs'], link=g['link'], bom=key))
        elif p[0] == HAVE:
            _, what, val, fp, note = p
            have.append(dict(what=what, val=val, fp=fp, note=note, refs=g['refs'], bom=key))
        else:
            na.append(dict(note=p[1], refs=g['refs'], bom=key))

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HEAD = PatternFill('solid', fgColor='404040')
    THIN = Border(bottom=Side('thin', color='D0D0D0'))

    def setup(ws, head, widths, wrap_cols, center_cols):
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = HEAD
            c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 30
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws._wrap, ws._center = wrap_cols, center_cols

    def putrow(ws, vals):
        ws.append(vals)
        i = ws.max_row
        for col in range(1, len(vals) + 1):
            cell = ws.cell(row=i, column=col)
            cell.border = THIN
            if col in ws._wrap:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            elif col in ws._center:
                cell.alignment = Alignment(horizontal='center', vertical='top')
            else:
                cell.alignment = Alignment(vertical='top')
        return i

    wb = openpyxl.Workbook()

    # ---------- лист 1: к заказу ----------
    ws = wb.active
    ws.title = 'К заказу'
    setup(ws, ['Что заказать', 'Номинал', 'Корпус', 'Заказать, шт', 'Нужно, шт',
               'Ссылка', 'Позиции', 'Плата', 'Примечание'],
          [30, 22, 22, 12, 10, 46, 22, 14, 70], wrap_cols={7, 9}, center_cols={4, 5})
    buy.sort(key=lambda r: (0 if r['what'].startswith(('10 нФ', 'CC', 'GRM', 'RC')) else 1, r['fp'], r['val']))
    warn = PatternFill('solid', fgColor='FCE4E4')
    for r in buy:
        refs, sheets = fmt(r['refs'], by_ref)
        i = putrow(ws, [r['what'], r['val'], r['fp'], r['oq'], r['need'],
                        r['link'] or 'ссылки в схеме нет', refs, sheets, r['note']])
        if r['link']:
            c = ws.cell(row=i, column=6)
            c.hyperlink = r['link']
            c.font = Font(color='0563C1', underline='single')
        else:
            ws.cell(row=i, column=6).font = Font(color='C00000', italic=True)
        if 'ВНИМАНИЕ' in r['note'] or 'НЕ брать' in r['note'] or 'ССЫЛКИ НЕТ' in r['note']:
            for col in range(1, 10):
                ws.cell(row=i, column=col).fill = warn
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:I%d' % ws.max_row
    ws.append([])
    ws.append(['ИТОГО %d позиций' % len(buy), '', '', sum(r['oq'] for r in buy),
               sum(r['need'] for r in buy)])
    for col in (1, 4, 5):
        ws.cell(row=ws.max_row, column=col).font = Font(bold=True)

    # ---------- лист 2: из запасов ----------
    ws2 = wb.create_sheet('Из запасов')
    setup(ws2, ['Что ставим', 'Номинал (фактический)', 'Корпус (фактический)', 'Кол-во',
                'Позиции', 'Плата', 'Вместо (PN в BOM)', 'Примечание'],
          [34, 26, 24, 9, 30, 14, 24, 74], wrap_cols={5, 8}, center_cols={4})
    have.sort(key=lambda r: (r['fp'], r['val']))
    for r in have:
        refs, sheets = fmt(r['refs'], by_ref)
        i = putrow(ws2, [r['what'], r['val'], r['fp'], len(r['refs']), refs, sheets,
                         r['bom'], r['note']])
        if 'ВНИМАНИЕ' in r['note']:
            for col in range(1, 9):
                ws2.cell(row=i, column=col).fill = warn
    for r in na:
        refs, sheets = fmt(r['refs'], by_ref)
        putrow(ws2, ['—', '', '', len(r['refs']), refs, sheets, r['bom'], r['note']])
    ws2.freeze_panes = 'A2'
    ws2.auto_filter.ref = 'A1:H%d' % ws2.max_row
    ws2.append([])
    ws2.append(['ИТОГО %d позиций' % (len(have) + len(na)), '', '',
                sum(len(r['refs']) for r in have) + sum(len(r['refs']) for r in na)])
    for col in (1, 4):
        ws2.cell(row=ws2.max_row, column=col).font = Font(bold=True)

    out = OUT
    if os.path.exists(OUT) and '--force' not in sys.argv:
        out = OUT.replace('.xlsx', '_new.xlsx')
        print('ВНИМАНИЕ: %s уже есть (в нём могут быть ваши пометки) — пишу рядом.'
              % os.path.basename(OUT))
        print('          перезаписать поверх: python stock/make_bom_xlsx.py --force')
    wb.save(out)
    n_buy = sum(r['need'] for r in buy)
    n_have = sum(len(r['refs']) for r in have) + sum(len(r['refs']) for r in na)
    print('сохранено: %s' % out)
    print('К заказу:   %2d позиций, %3d шт на платы, %3d шт с запасом'
          % (len(buy), n_buy, sum(r['oq'] for r in buy)))
    print('Из запасов: %2d позиций, %3d шт' % (len(have) + len(na), n_have))
    print('всего компонентов %d (в нетлисте %d) %s'
          % (n_buy + n_have, len(comps), 'OK' if n_buy + n_have == len(comps) else 'РАСХОЖДЕНИЕ!'))
    if unknown:
        print('НЕТ В PLAN: %s' % ', '.join(unknown))
    return 1 if unknown or n_buy + n_have != len(comps) else 0


if __name__ == '__main__':
    raise SystemExit(main())
